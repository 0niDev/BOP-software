"""Report export utilities - PDF, Excel, CSV, Print."""
from __future__ import annotations

import csv
import os
from datetime import datetime
from io import BytesIO
from typing import Any

from PySide6.QtCore import Qt, QMarginsF
from PySide6.QtGui import QTextDocument, QPageLayout, QPageSize
from PySide6.QtPrintSupport import QPrinter, QPrintDialog
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox, QTextEdit, QDialog, QVBoxLayout, QHBoxLayout, QPushButton, QLabel

from utils.logger import get_logger

logger = get_logger(__name__)


class ReportExporter:
    """Export reports to various formats."""

    @staticmethod
    def print_report(text_edit: QTextEdit, title: str = "Report") -> bool:
        """Print the report directly."""
        try:
            printer = QPrinter(QPrinter.HighResolution)
            
            # Create page layout with proper margins
            page_layout = QPageLayout(
                QPageSize(QPageSize.A4),
                QPageLayout.Orientation.Portrait,
                QMarginsF(10, 10, 10, 10),  # margins in mm
                QPageLayout.Unit.Millimeter
            )
            printer.setPageLayout(page_layout)
            printer.setDocName(title)
            printer.setCreator("Pharmaceutical ERP")
            
            dialog = QPrintDialog(printer)
            if dialog.exec() != QPrintDialog.Accepted:
                return False
            
            doc = QTextDocument()
            doc.setHtml(text_edit.toHtml())
            doc.print_(printer)
            return True
            
        except Exception as e:
            logger.exception(f"Print failed: {e}")
            return False

    @staticmethod
    def export_pdf(text_edit: QTextEdit, default_filename: str = "report.pdf") -> bool:
        """Export report to PDF."""
        try:
            file_path, _ = QFileDialog.getSaveFileName(
                None,
                "Save PDF",
                default_filename,
                "PDF Files (*.pdf)"
            )
            
            if not file_path:
                return False
            
            if not file_path.endswith('.pdf'):
                file_path += '.pdf'
            
            printer = QPrinter(QPrinter.HighResolution)
            printer.setOutputFormat(QPrinter.PdfFormat)
            printer.setOutputFileName(file_path)
            
            # Create page layout with proper margins
            page_layout = QPageLayout(
                QPageSize(QPageSize.A4),
                QPageLayout.Orientation.Portrait,
                QMarginsF(10, 10, 10, 10),  # margins in mm
                QPageLayout.Unit.Millimeter
            )
            printer.setPageLayout(page_layout)
            printer.setDocName(default_filename)
            
            doc = QTextDocument()
            doc.setHtml(text_edit.toHtml())
            doc.print_(printer)
            
            QMessageBox.information(None, "Success", f"PDF exported to:\n{file_path}")
            logger.info(f"PDF exported: {file_path}")
            return True
            
        except Exception as e:
            logger.exception(f"PDF export failed: {e}")
            QMessageBox.warning(None, "Export Failed", f"Failed to export PDF:\n{str(e)}")
            return False

    @staticmethod
    def export_excel(text_edit: QTextEdit, default_filename: str = "report.xlsx") -> bool:
        """Export report to Excel."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment
            
            file_path, _ = QFileDialog.getSaveFileName(
                None,
                "Save Excel",
                default_filename,
                "Excel Files (*.xlsx)"
            )
            
            if not file_path:
                return False
            
            if not file_path.endswith('.xlsx'):
                file_path += '.xlsx'
            
            html = text_edit.toHtml()
            
            wb = Workbook()
            ws = wb.active
            
            data = ReportExporter._parse_html_table(html)
            
            if not data:
                data = ReportExporter._parse_html_text(html)
            
            for row_idx, row in enumerate(data, 1):
                for col_idx, value in enumerate(row, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    if row_idx == 1:
                        cell.font = Font(bold=True)
                        cell.alignment = Alignment(horizontal='center')
            
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column].width = adjusted_width
            
            wb.save(file_path)
            
            QMessageBox.information(None, "Success", f"Excel exported to:\n{file_path}")
            logger.info(f"Excel exported: {file_path}")
            return True
            
        except ImportError:
            QMessageBox.warning(None, "Missing Library", 
                "openpyxl is not installed.\nPlease install it with:\npip install openpyxl")
            return False
        except Exception as e:
            logger.exception(f"Excel export failed: {e}")
            QMessageBox.warning(None, "Export Failed", f"Failed to export Excel:\n{str(e)}")
            return False

    @staticmethod
    def export_csv(text_edit: QTextEdit, default_filename: str = "report.csv") -> bool:
        """Export report to CSV."""
        try:
            file_path, _ = QFileDialog.getSaveFileName(
                None,
                "Save CSV",
                default_filename,
                "CSV Files (*.csv)"
            )
            
            if not file_path:
                return False
            
            if not file_path.endswith('.csv'):
                file_path += '.csv'
            
            html = text_edit.toHtml()
            data = ReportExporter._parse_html_table(html)
            
            if not data:
                data = ReportExporter._parse_html_text(html)
            
            with open(file_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                for row in data:
                    writer.writerow(row)
            
            QMessageBox.information(None, "Success", f"CSV exported to:\n{file_path}")
            logger.info(f"CSV exported: {file_path}")
            return True
            
        except Exception as e:
            logger.exception(f"CSV export failed: {e}")
            QMessageBox.warning(None, "Export Failed", f"Failed to export CSV:\n{str(e)}")
            return False

    @staticmethod
    def _parse_html_table(html: str) -> list[list]:
        """Parse HTML table into 2D list."""
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(html, 'html.parser')
            data = []
            
            tables = soup.find_all('table')
            if not tables:
                return []
            
            for table in tables:
                for row in table.find_all('tr'):
                    row_data = []
                    for cell in row.find_all(['td', 'th']):
                        text = cell.get_text(strip=True)
                        # Try to convert to number
                        try:
                            cleaned = text.replace('Rs.', '').replace(',', '').strip()
                            if cleaned and cleaned.replace('.', '').replace('-', '').isdigit():
                                row_data.append(float(cleaned))
                            else:
                                row_data.append(text)
                        except:
                            row_data.append(text)
                    if row_data:
                        data.append(row_data)
            
            return data
        except Exception as e:
            logger.warning(f"HTML table parsing failed: {e}")
            return []

    @staticmethod
    def _parse_html_text(html: str) -> list[list]:
        """Parse HTML text into 2D list."""
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(html, 'html.parser')
            text = soup.get_text()
            lines = text.strip().split('\n')
            data = []
            for line in lines:
                line = line.strip()
                if line and not line.startswith('=') and not line.startswith('-'):
                    import re
                    parts = re.split(r'\s{2,}|\t', line)
                    if len(parts) > 1:
                        data.append(parts)
                    else:
                        data.append([line])
            return data
        except Exception as e:
            logger.warning(f"HTML text parsing failed: {e}")
            return []

    @staticmethod
    def show_export_dialog(parent, text_edit: QTextEdit, report_name: str = "Report"):
        """Show a dialog with export options."""
        if not text_edit or not text_edit.toPlainText().strip():
            QMessageBox.information(parent, "No Data", "Please generate the report first.")
            return
        
        dialog = QDialog(parent)
        dialog.setWindowTitle("Export Report")
        dialog.setModal(True)
        dialog.resize(450, 250)
        
        layout = QVBoxLayout(dialog)
        
        label = QLabel(f"Export '{report_name.replace('_', ' ')}' to:")
        label.setStyleSheet("font-size: 14px; font-weight: bold;")
        layout.addWidget(label)
        
        btn_layout = QHBoxLayout()
        
        print_btn = QPushButton("🖨️ Print")
        print_btn.clicked.connect(lambda: ReportExporter.print_report(text_edit, report_name))
        print_btn.setMinimumHeight(50)
        print_btn.setMinimumWidth(80)
        btn_layout.addWidget(print_btn)
        
        pdf_btn = QPushButton("📄 PDF")
        pdf_btn.clicked.connect(lambda: ReportExporter.export_pdf(text_edit, f"{report_name}.pdf"))
        pdf_btn.setMinimumHeight(50)
        pdf_btn.setMinimumWidth(80)
        btn_layout.addWidget(pdf_btn)
        
        excel_btn = QPushButton("📊 Excel")
        excel_btn.clicked.connect(lambda: ReportExporter.export_excel(text_edit, f"{report_name}.xlsx"))
        excel_btn.setMinimumHeight(50)
        excel_btn.setMinimumWidth(80)
        btn_layout.addWidget(excel_btn)
        
        csv_btn = QPushButton("📋 CSV")
        csv_btn.clicked.connect(lambda: ReportExporter.export_csv(text_edit, f"{report_name}.csv"))
        csv_btn.setMinimumHeight(50)
        csv_btn.setMinimumWidth(80)
        btn_layout.addWidget(csv_btn)
        
        layout.addLayout(btn_layout)
        
        close_btn = QPushButton("Close")
        close_btn.clicked.connect(dialog.accept)
        close_btn.setMinimumHeight(30)
        layout.addWidget(close_btn)
        
        dialog.exec()