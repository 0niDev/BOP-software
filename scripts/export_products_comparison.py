"""Export a comparison Excel of all products from PharmaPro Full Export vs current DB.

Output: EXPORTS/Products_Comparison.xlsx
Sheets: Summary, Matched, Only in Old Export, Only in Current DB, All Products
"""
from __future__ import annotations

import csv
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from utils.env_loader import setup_import_env
setup_import_env()

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from database.connection import get_db
EXPORT_DIR = os.path.join(ROOT, "PharmaPro_FullExport")
OUT_DIR = os.path.join(ROOT, "EXPORTS")

TYPE_LABEL = {
    "00": "Finished Good",
    "01": "Raw Material",
    "02": "Packing Material",
    "03": "Semi-Finished (Bulk)",
}

TITLE_FILL = PatternFill("solid", fgColor="1F3864")
TITLE_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")


def load_old_products() -> dict:
    path = os.path.join(EXPORT_DIR, "Products.csv")
    products = {}
    with open(path, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            pid = r["ProductId"].strip().zfill(5)
            name = (r.get("ProductName") or "").strip()
            if not name:
                continue
            cid = (r.get("CompanyId") or "").strip()
            products[pid] = {
                "id": pid,
                "name": name,
                "company_id": cid,
                "type": TYPE_LABEL.get(cid, "Other"),
                "packing": (r.get("Packing") or "").strip(),
                "purchase_price": r.get("PurchasePrice", "0"),
                "trade_price": r.get("TradePrice", "0"),
                "retail_price": r.get("RetailPrice", "0"),
                "min_stock": r.get("MinStockLevel", "0"),
                "is_active": r.get("IsActive", "True"),
            }
    return products


def load_new_products() -> dict:
    db = get_db()
    rows = db.fetch_all(
        "SELECT id, item_code, item_name, item_type, unit, "
        "purchase_price, selling_price, is_active FROM items"
    )
    products = {}
    for row in rows:
        code = (row["item_code"] or "").strip()
        name = (row["item_name"] or "").strip()
        if not name:
            continue
        products[code] = {
            "id": row["id"],
            "code": code,
            "name": name,
            "item_type": row["item_type"] or "",
            "unit": row["unit"] or "",
            "purchase_price": row["purchase_price"] or 0,
            "selling_price": row["selling_price"] or 0,
            "is_active": row["is_active"],
        }
    return products


def write_header(ws, headers, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col, text in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=text)
        c.font = TITLE_FONT
        c.fill = TITLE_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def write_row(ws, row_num, values, fill=None):
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row_num, column=col, value=v)
        c.border = BORDER
        if fill:
            c.fill = fill


def main():
    print("Loading old export products...")
    old_products = load_old_products()
    print(f"  {len(old_products)} products")

    print("Loading current DB products...")
    new_products = load_new_products()
    print(f"  {len(new_products)} products")

    old_by_name = {p["name"].strip().lower(): (pid, p) for pid, p in old_products.items()}
    new_by_name = {p["name"].strip().lower(): (code, p) for code, p in new_products.items()}

    matched, only_old, only_new = [], [], []
    for name_lower, (pid, op) in old_by_name.items():
        if name_lower in new_by_name:
            code, np = new_by_name[name_lower]
            matched.append((pid, op, code, np))
        else:
            only_old.append((pid, op))
    for name_lower, (code, np) in new_by_name.items():
        if name_lower not in old_by_name:
            only_new.append((code, np))

    print(f"  Matched: {len(matched)}, Only old: {len(only_old)}, Only new: {len(only_new)}")

    wb = Workbook()

    # ---- Sheet 1: Summary ----
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 15
    ws.cell(row=1, column=1, value="Products Comparison Summary").font = Font(bold=True, size=14)
    ws.merge_cells("A1:B1")

    summary = [
        ("Total Old Export Products", len(old_products)),
        ("Total Current DB Products", len(new_products)),
        ("Matched (same name)", len(matched)),
        ("Only in Old Export", len(only_old)),
        ("Only in Current DB", len(only_new)),
    ]
    for i, (label, val) in enumerate(summary, 3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)

    ws.cell(row=9, column=1, value="Old Export by Type").font = Font(bold=True)
    type_counts = {}
    for p in old_products.values():
        type_counts[p["type"]] = type_counts.get(p["type"], 0) + 1
    for i, (t, c) in enumerate(sorted(type_counts.items()), 10):
        ws.cell(row=i, column=1, value=t)
        ws.cell(row=i, column=2, value=c)

    ws.cell(row=15, column=1, value="Current DB by Type").font = Font(bold=True)
    type_counts2 = {}
    for p in new_products.values():
        type_counts2[p["item_type"]] = type_counts2.get(p["item_type"], 0) + 1
    for i, (t, c) in enumerate(sorted(type_counts2.items()), 16):
        ws.cell(row=i, column=1, value=t)
        ws.cell(row=i, column=2, value=c)

    # ---- Sheet 2: Matched ----
    ws2 = wb.create_sheet("Matched")
    h = ["Old ID", "Product Name", "Old Type", "Old Packing", "Old Purchase Price",
         "New Code", "New Item Type", "New Unit", "New Purchase Price", "New Selling Price"]
    w = [10, 42, 18, 12, 16, 12, 18, 10, 16, 16]
    write_header(ws2, h, w)
    for i, (pid, op, code, np) in enumerate(matched, 2):
        write_row(ws2, i, [pid, op["name"], op["type"], op["packing"], op["purchase_price"],
                           code, np["item_type"], np["unit"], np["purchase_price"],
                           np["selling_price"]], GREEN_FILL)
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(h))}{len(matched)+1}"

    # ---- Sheet 3: Only in Old ----
    ws3 = wb.create_sheet("Only in Old Export")
    h = ["Old ID", "Product Name", "Type", "Packing", "Purchase Price",
         "Trade Price", "Retail Price", "Min Stock", "Active"]
    w = [10, 42, 18, 12, 16, 16, 16, 12, 10]
    write_header(ws3, h, w)
    for i, (pid, op) in enumerate(only_old, 2):
        write_row(ws3, i, [pid, op["name"], op["type"], op["packing"],
                           op["purchase_price"], op["trade_price"], op["retail_price"],
                           op["min_stock"], op["is_active"]], RED_FILL)
    ws3.auto_filter.ref = f"A1:{get_column_letter(len(h))}{len(only_old)+1}"

    # ---- Sheet 4: Only in New ----
    ws4 = wb.create_sheet("Only in Current DB")
    h = ["Code", "Product Name", "Item Type", "Unit", "Purchase Price",
         "Selling Price", "Active"]
    w = [12, 42, 18, 10, 16, 16, 10]
    write_header(ws4, h, w)
    for i, (code, np) in enumerate(only_new, 2):
        write_row(ws4, i, [code, np["name"], np["item_type"], np["unit"],
                           np["purchase_price"], np["selling_price"],
                           "Yes" if np["is_active"] else "No"], YELLOW_FILL)
    ws4.auto_filter.ref = f"A1:{get_column_letter(len(h))}{len(only_new)+1}"

    # ---- Sheet 5: All Products ----
    ws5 = wb.create_sheet("All Products")
    h = ["Status", "Old ID", "New Code", "Product Name", "Type", "Packing/Unit",
         "Old Purchase Price", "New Purchase Price", "New Selling Price"]
    w = [14, 10, 12, 42, 18, 14, 18, 18, 18]
    write_header(ws5, h, w)
    row = 2
    for pid, op, code, np in matched:
        write_row(ws5, row, ["MATCHED", pid, code, op["name"], op["type"],
                             op["packing"] or np["unit"], op["purchase_price"],
                             np["purchase_price"], np["selling_price"]], GREEN_FILL)
        row += 1
    for pid, op in only_old:
        write_row(ws5, row, ["OLD ONLY", pid, "", op["name"], op["type"], op["packing"],
                             op["purchase_price"], "", ""], RED_FILL)
        row += 1
    for code, np in only_new:
        write_row(ws5, row, ["NEW ONLY", "", code, np["name"], np["item_type"], np["unit"],
                             "", np["purchase_price"], np["selling_price"]], YELLOW_FILL)
        row += 1
    ws5.auto_filter.ref = f"A1:{get_column_letter(len(h))}{row-1}"

    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, "Products_Comparison.xlsx")
    wb.save(out)
    print(f"\nSaved: {out}")


if __name__ == "__main__":
    main()
