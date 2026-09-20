"""Export PharmaPro formulations (formula + filling) into a readable Excel workbook.

Data sources: ``PharmaPro_FullExport``
    Products.csv        product master
                        CompanyId: 00 = finished goods, 01 = raw materials,
                                   02 = packing materials, 03 = bulk / semi-finished
    FormulaHeader.csv / FormulaBody.csv
                        bulk (semi-finished) master recipes.
                        Header product (03/02) <- components (01 raw materials)
    FillingHeader.csv / FillingBody.csv
                        filling / packing recipes.
                        Header product (00 finished good) <- components
                        (02 packing materials + 03 bulk to be filled)

A finished good is produced by filling a bulk, so the complete formulation of a
finished good is the bulk master formula expanded through the filling quantity,
plus the packing materials consumed at filling time.

Output: ``EXPORTS/Formulations.xlsx``

Sheets
    Formulations      one block per product: heading (product name), then the
                      RAW MATERIALS list, then the PACKING MATERIALS list,
                      then the next product
    All Components    flat "full sheet": every product / component pair
    Products Index    every product with component counts and status

Run: python scripts/export_formulations.py
"""
from __future__ import annotations

import csv
import os
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXPORT_DIR = os.path.join(ROOT, "PharmaPro_FullExport")
OUT_DIR = os.path.join(ROOT, "EXPORTS")

# CompanyId -> human label
COMPANY_LABEL = {
    "00": "Finished Good",
    "01": "Raw Material",
    "02": "Packing Material",
    "03": "Semi-Finished (Bulk)",
}

# ---------------------------------------------------------------- styling ---
TITLE_FILL = PatternFill("solid", fgColor="1F3864")
TITLE_FONT = Font(bold=True, size=12, color="FFFFFF")
INFO_FONT = Font(italic=True, size=9, color="404040")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
SECTION_FONT = Font(bold=True, size=10, color="1F3864")
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
HEAD_FONT = Font(bold=True, size=10)
NOTE_FONT = Font(italic=True, size=9, color="C00000")
THIN = Side(style="thin", color="BFBFBF")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
QTY_FMT = "0.##########"


def load_csv(name: str) -> list:
    path = os.path.join(EXPORT_DIR, name)
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        return [{k.strip("\ufeff"): v for k, v in r.items()} for r in csv.DictReader(fh)]


def code(value: str) -> str:
    """Normalise a PharmaPro product code (some exports drop leading zeros)."""
    value = (value or "").strip()
    if not value or not value.isdigit():
        return value
    return value.zfill(5)


def fmt_qty(value: float) -> str:
    """Compact human readable quantity (1.0 -> '1', 0.0015 -> '0.0015')."""
    text = f"{value:.10f}".rstrip("0").rstrip(".")
    return text or "0"


class FormulationData:
    """Loads the export and resolves formula -> filling -> finished-good chain."""

    def __init__(self) -> None:
        self.products = {}
        for row in load_csv("Products.csv"):
            pid = code(row.get("ProductId", ""))
            row["ProductId"] = pid
            self.products[pid] = row

        self.warnings = []

        # ---- bulk master formulas: product -> {component: qty_per_unit} -----
        formula_qty = defaultdict(lambda: defaultdict(float))
        for row in load_csv("FormulaBody.csv"):
            fid = (row.get("FormulaID") or "").strip()
            cid = code(row.get("ProductID", ""))
            if not fid or not cid:
                continue
            formula_qty[fid][cid] += self._to_float(row.get("Qty"))

        self.formula_of = {}
        self.formula_head_products = set()
        for row in load_csv("FormulaHeader.csv"):
            fid = (row.get("FormulaID") or "").strip()
            pid = code(row.get("ProductID", ""))
            if not fid or not pid:
                continue
            self.formula_head_products.add(pid)
            base = self._to_float(row.get("Qty")) or 1.0
            for cid, qty in formula_qty.get(fid, {}).items():
                self.formula_of.setdefault(pid, {})[cid] = qty / base

        # ---- filling / packing formulas: finished good -> {component: qty} --
        filling_qty = defaultdict(lambda: defaultdict(float))
        for row in load_csv("FillingBody.csv"):
            fid = (row.get("FormulaID") or "").strip()
            cid = code(row.get("ProductID", ""))
            if not fid or not cid:
                continue
            filling_qty[fid][cid] += self._to_float(row.get("Qty"))

        self.filling_of = {}
        self.filling_head_products = set()
        for row in load_csv("FillingHeader.csv"):
            fid = (row.get("FormulaID") or "").strip()
            pid = code(row.get("ProductID", ""))
            if not fid or not pid:
                continue
            self.filling_head_products.add(pid)
            base = self._to_float(row.get("Qty")) or 1.0
            for cid, qty in filling_qty.get(fid, {}).items():
                self.filling_of.setdefault(pid, {})[cid] = qty / base

        # products referenced by a formula but absent from the master
        referenced = set(self.formula_of) | set(self.filling_of)
        for pid in sorted(referenced):
            referenced.update(self.formula_of.get(pid, {}))
            referenced.update(self.filling_of.get(pid, {}))
        for pid in sorted(referenced - set(self.products)):
            self.warnings.append(
                f"{pid} referenced by a formula but missing from Products.csv"
            )

    @staticmethod
    def _to_float(value) -> float:
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    # ---------------------------------------------------------------- helpers
    def name(self, pid: str) -> str:
        p = self.products.get(pid)
        if p and p.get("ProductName"):
            return p["ProductName"].strip()
        return f"(unknown {pid})"

    def company(self, pid: str) -> str:
        p = self.products.get(pid)
        return (p.get("CompanyId") or "??").strip() if p else "??"

    def ptype(self, pid: str) -> str:
        if pid not in self.products:
            return "Not in Master"
        return COMPANY_LABEL.get(self.company(pid), "Other")

    def pack(self, pid: str) -> str:
        p = self.products.get(pid)
        return (p.get("Packing") or "").strip() if p else ""

    def product_codes(self) -> list:
        """Every product of interest, finished goods first.

        Includes products that only have an (empty) formula/filling header so
        that missing recipes are reported instead of silently dropped, and
        every finished good so the sheet covers the whole product range.
        """
        ids = set(self.formula_of) | set(self.filling_of)
        ids |= self.formula_head_products | self.filling_head_products
        ids |= {pid for pid, row in self.products.items()
                if (row.get("CompanyId") or "").strip() == "00"}
        order = {"00": 0, "03": 1, "02": 2, "01": 3}

        def sort_key(pid: str):
            return (order.get(self.company(pid), 9), pid)

        return sorted(ids, key=sort_key)

    # ------------------------------------------------------------- the recipe
    def build(self, pid: str) -> dict:
        """Resolve one product into raw-material and packing-material lines."""
        raw = []            # (component, qty_per_unit, via_bulk_id)
        packing = []        # (component, qty_per_unit)
        bulks = []          # bulk ids filled into this product
        notes = []

        filling = self.filling_of.get(pid)
        if filling:
            for cid, qty in filling.items():
                comp_type = self.company(cid)
                if comp_type == "02":
                    packing.append((cid, qty))
                elif comp_type == "03":
                    bulks.append((cid, qty))
                    bulk_formula = self.formula_of.get(cid)
                    if not bulk_formula:
                        # keep the bulk visible with the quantity actually
                        # required, so no source row is silently dropped
                        notes.append(
                            f"no master formula on file for bulk {cid} {self.name(cid)}"
                        )
                        raw.append((cid, round(qty, 12), ""))
                        continue
                    for rcid, rqty in bulk_formula.items():
                        raw.append((rcid, round(qty * rqty, 12), cid))
                elif comp_type == "01":
                    raw.append((cid, qty, ""))
                else:
                    packing.append((cid, qty))
        elif pid in self.formula_of:
            # a bulk / packing item that carries its own master formula
            for cid, qty in self.formula_of[pid].items():
                if self.company(cid) == "02":
                    packing.append((cid, qty))
                else:
                    raw.append((cid, qty, ""))

        # ---- explain every recipe that could not be resolved ----------------
        is_finished = self.company(pid) == "00"
        if not filling and not raw and not packing:
            if pid in self.filling_head_products:
                notes.append("filling formula recorded but it has no component lines")
            elif pid in self.formula_head_products:
                notes.append("master formula recorded but it has no component lines")
        if is_finished and pid not in self.filling_head_products:
            notes.append("no filling / packing formula in the export")
        if not is_finished and self.company(pid) == "03" \
                and pid not in self.formula_head_products:
            notes.append("bulk has no master formula in the export")

        # merge duplicates created by expansion, keeping the bulk trail
        merged = {}
        for cid, qty, via in raw:
            merged[(cid, via)] = merged.get((cid, via), 0.0) + qty
        raw = sorted(
            ((cid, round(qty, 12), via) for (cid, via), qty in merged.items()),
            key=lambda line: line[0],
        )

        packed = {}
        for cid, qty in packing:
            packed[cid] = packed.get(cid, 0.0) + qty
        packing = sorted((cid, round(qty, 12)) for cid, qty in packed.items())

        bulk_map = {}
        for cid, qty in bulks:
            bulk_map[cid] = bulk_map.get(cid, 0.0) + qty

        return {
            "id": pid,
            "name": self.name(pid),
            "type": self.ptype(pid),
            "pack": self.pack(pid),
            "raw": raw,
            "packing": packing,
            "bulks": sorted(bulk_map.items()),
            "notes": notes,
        }


# ------------------------------------------------------------------ writers --
def write_formulations_sheet(ws, data: FormulationData, recipes: list) -> None:
    """One block per product: heading, raw materials, packing materials."""
    for i, width in enumerate([14, 46, 22, 18, 12, 34], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    row = 1
    for rec in recipes:
        # ---- product heading ("the name of the project") --------------------
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.fill = TITLE_FILL
            cell.font = TITLE_FONT
        ws.cell(row=row, column=1, value=f"{rec['id']}   {rec['name']}")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row=row, column=1).alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 20
        row += 1

        info = f"Type: {rec['type']}"
        if rec["pack"]:
            info += f"   |   Pack size: {rec['pack']}"
        if rec["bulks"]:
            info += "   |   Bulk filled: " + ", ".join(
                f"{b} {data.name(b)} x {fmt_qty(q)}" for b, q in rec["bulks"]
            )
        info += f"   |   Quantity basis: 1 unit of {rec['id']}"
        ws.cell(row=row, column=1, value=info).font = INFO_FONT
        row += 1

        if rec["notes"]:
            ws.cell(row=row, column=1, value="Note: " + "; ".join(rec["notes"])).font = NOTE_FONT
            row += 1

        for label, lines, is_packing in (
            ("RAW MATERIALS", rec["raw"], False),
            ("PACKING MATERIALS", rec["packing"], True),
        ):
            for col in range(1, 7):
                cell = ws.cell(row=row, column=col)
                cell.fill = SECTION_FILL
                cell.font = SECTION_FONT
            ws.cell(row=row, column=1, value=label)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            row += 1

            if not lines:
                ws.cell(row=row, column=1, value="(none defined)").font = INFO_FONT
                row += 1
                continue

            for col, text in enumerate(
                ["Component ID", "Component Name", "Component Type", "Qty", "UoM", "Source"],
                start=1,
            ):
                cell = ws.cell(row=row, column=col, value=text)
                cell.font = HEAD_FONT
                cell.fill = HEAD_FILL
                cell.border = CELL_BORDER
            row += 1

            for line in lines:
                if is_packing:
                    cid, qty = line
                    via = ""
                else:
                    cid, qty, via = line
                if via:
                    source = f"bulk {via} {data.name(via)}"
                elif is_packing:
                    source = "filling formula"
                elif data.company(cid) == "03":
                    source = "unexpanded bulk (master formula missing)"
                else:
                    source = "master formula"
                values = [cid, data.name(cid), data.ptype(cid), qty, data.pack(cid), source]
                for col, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.border = CELL_BORDER
                    if col == 4:
                        cell.number_format = QTY_FMT
                row += 1

        row += 1  # blank separator before the next product

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def _write_header_row(ws, headers: list, widths: list) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = TITLE_FONT
        cell.fill = TITLE_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def _write_flat_row(ws, row: int, values: list) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        if col == 8:
            cell.number_format = QTY_FMT


def write_all_components_sheet(ws, data: FormulationData, recipes: list) -> None:
    """Full flat sheet: every product against every one of its components."""
    headers = [
        "Product ID", "Product Name", "Product Type", "Pack Size",
        "Component ID", "Component Name", "Component Type",
        "Qty per 1 Unit", "UoM", "Component Group", "Via Bulk",
    ]
    _write_header_row(ws, headers, [12, 42, 21, 12, 13, 42, 18, 16, 11, 17, 34])

    row = 2
    for rec in recipes:
        for cid, qty, via in rec["raw"]:
            values = [
                rec["id"], rec["name"], rec["type"], rec["pack"],
                cid, data.name(cid), data.ptype(cid), qty, data.pack(cid),
                data.ptype(cid),
                f"{via} {data.name(via)}".strip() if via else "",
            ]
            _write_flat_row(ws, row, values)
            row += 1

        for cid, qty in rec["packing"]:
            values = [
                rec["id"], rec["name"], rec["type"], rec["pack"],
                cid, data.name(cid), data.ptype(cid), qty, data.pack(cid),
                data.ptype(cid), "",
            ]
            _write_flat_row(ws, row, values)
            row += 1

    ws.auto_filter.ref = f"A1:K{row - 1}"


def write_index_sheet(ws, data: FormulationData, recipes: list) -> None:
    """Every product in the master, with component counts and status."""
    headers = [
        "Product ID", "Product Name", "Product Type", "Pack Size",
        "Raw Materials", "Packing Materials", "Total Components",
        "Has Formulation", "Company", "Notes",
    ]
    _write_header_row(ws, headers, [12, 46, 21, 12, 14, 17, 17, 15, 10, 60])

    by_id = {rec["id"]: rec for rec in recipes}
    # master products first, then any formula-only product missing from the master
    orphans = sorted(set(by_id) - set(data.products))
    listed = sorted(data.products) + orphans
    row = 2
    for pid in listed:
        rec = by_id.get(pid)
        in_master = pid in data.products
        n_raw = len(rec["raw"]) if rec else 0
        n_pack = len(rec["packing"]) if rec else 0
        notes = "; ".join(rec["notes"]) if rec and rec["notes"] else ""
        if not in_master:
            notes = "formula record exists but the product is missing from Products.csv"
        elif not notes and not (n_raw + n_pack):
            notes = "no formula / filling record in the export"
        values = [
            pid,
            data.name(pid),
            data.ptype(pid),
            data.pack(pid),
            n_raw,
            n_pack,
            n_raw + n_pack,
            "Yes" if (n_raw + n_pack) else "No",
            data.company(pid) if in_master else "-",
            notes,
        ]
        _write_flat_row(ws, row, values)
        row += 1

    ws.auto_filter.ref = f"A1:J{row - 1}"


def main() -> None:
    t0 = datetime.now()
    data = FormulationData()
    print(f"Loaded {len(data.products)} products, {len(data.formula_of)} master formulas, "
          f"{len(data.filling_of)} filling formulas")

    recipes = []
    for pid in data.product_codes():
        rec = data.build(pid)
        if rec["raw"] or rec["packing"] or rec["notes"]:
            recipes.append(rec)

    total_raw = sum(len(r["raw"]) for r in recipes)
    total_pack = sum(len(r["packing"]) for r in recipes)
    empty = [r for r in recipes if not r["raw"] and not r["packing"]]
    print(f"Resolved {len(recipes)} product blocks -> {total_raw} raw-material lines, "
          f"{total_pack} packing-material lines")
    print(f"{len(recipes) - len(empty)} products have a usable formulation; "
          f"{len(empty)} are listed as unresolved (see the Notes column)")

    wb = Workbook()
    write_formulations_sheet(wb.active, data, recipes)
    wb.active.title = "Formulations"
    write_all_components_sheet(wb.create_sheet("All Components"), data, recipes)
    write_index_sheet(wb.create_sheet("Products Index"), data, recipes)

    os.makedirs(OUT_DIR, exist_ok=True)
    out_path = os.path.join(OUT_DIR, "Formulations.xlsx")
    wb.save(out_path)

    for warning in data.warnings:
        print("WARN:", warning)
    print(f"Saved: {out_path}")
    print(f"Elapsed: {(datetime.now() - t0).total_seconds():.1f}s")


if __name__ == "__main__":
    main()